# -*- coding: utf-8 -*-
"""Expense statistics aggregation (module 4)."""
from collections import defaultdict

from backend.extensions import db
from backend.models import StationMaster
from backend.util.helpers import is_empty, norm, to_float


def _parse_ratio(value):
    """Parse a share-ratio field like '30%' or 0.3 into a float fraction."""
    if is_empty(value):
        return 0.0
    s = str(value).strip().replace("%", "")
    f = to_float(s)
    if f is None:
        return 0.0
    return f / 100.0 if f > 1 else f


def _apply_filters(q, filters):
    city = norm(filters.get("city"))
    if city:
        q = q.filter(StationMaster.city == city)
    prop = norm(filters.get("property"))
    if prop:
        q = q.filter(StationMaster.property_name == prop)
    method = norm(filters.get("settle_method"))
    if method:
        q = q.filter(StationMaster.settle_method == method)
    need = norm(filters.get("need_settle"))
    if need:
        q = q.filter(StationMaster.need_settle == need)
    paid = norm(filters.get("is_paid"))
    if paid:
        q = q.filter(StationMaster.is_paid == paid)
    return q


def _decompose(st):
    """Break a station's expenses into category buckets."""
    settle = to_float(st.settle_amount) or 0.0
    elec = settle if (norm(st.need_settle) == "是" and norm(st.settle_method) != "不结算") else 0.0
    meter = to_float(st.last_meter_reading) or 0.0
    svc_price = to_float(st.service_price) or 0.0
    service = (svc_price * meter) if (svc_price and meter) else 0.0
    site_fee = to_float(st.annual_site_fee) or 0.0
    deposit = to_float(st.deposit) or 0.0
    share = _parse_ratio(st.share_ratio) * settle if settle else 0.0
    return {
        "电费": elec,
        "服务费": service,
        "场地费": site_fee,
        "押金": deposit,
        "分成": share,
    }


def build_report(filters=None):
    """Build the full expense report dict for the dashboard + export."""
    filters = filters or {}
    q = _apply_filters(StationMaster.query, filters)
    stations = q.all()

    cat_totals = defaultdict(float)
    by_property = defaultdict(float)
    by_city = defaultdict(float)
    trend = defaultdict(float)
    count = 0

    for st in stations:
        parts = _decompose(st)
        total = sum(parts.values())
        if total <= 0:
            continue
        count += 1
        for k, v in parts.items():
            cat_totals[k] += v
        prop = norm(st.property_name) or "未知物业"
        by_property[prop] += total
        city = norm(st.city) or "未知城市"
        by_city[city] += total
        month = (norm(st.next_settle_date) or "")[:7]
        if month:
            trend[month] += total

    grand = sum(cat_totals.values())
    metrics = {
        "total": round(grand, 2),
        "电费": round(cat_totals.get("电费", 0.0), 2),
        "服务费": round(cat_totals.get("服务费", 0.0), 2),
        "场地费": round(cat_totals.get("场地费", 0.0), 2),
        "押金": round(cat_totals.get("押金", 0.0), 2),
        "分成": round(cat_totals.get("分成", 0.0), 2),
        "site_count": count,
        "环比": None,  # no historical baseline available
    }

    return {
        "metrics": metrics,
        "categories": [
            {"name": k, "value": round(v, 2)}
            for k, v in sorted(cat_totals.items(), key=lambda x: -x[1])
        ],
        "by_property": [
            {"name": k, "value": round(v, 2)}
            for k, v in sorted(by_property.items(), key=lambda x: -x[1])[:20]
        ],
        "by_city": [
            {"name": k, "value": round(v, 2)}
            for k, v in sorted(by_city.items(), key=lambda x: -x[1])
        ],
        "trend": [
            {"month": k, "value": round(v, 2)}
            for k, v in sorted(trend.items())
        ],
        "filters": {k: norm(v) for k, v in filters.items() if norm(v)},
    }


def export_workbook(filters=None):
    """Build an openpyxl Workbook for the expense report. Returns Workbook."""
    import openpyxl
    from openpyxl.styles import Font

    report = build_report(filters)
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "指标卡"
    ws.append(["指标", "金额(元)"])
    for c in ("total", "电费", "服务费", "场地费", "押金", "分成"):
        label = {"total": "总支出", "电费": "电费", "服务费": "服务费",
                 "场地费": "场地费", "押金": "押金", "分成": "分成"}[c]
        ws.append([label, report["metrics"].get(c)])
    ws.append(["网点数", report["metrics"]["site_count"]])

    def dump(title, rows, name_key="name"):
        s = wb.create_sheet(title)
        s.append(["维度", "金额(元)"])
        for r in rows:
            s.append([r.get(name_key), r.get("value")])

    dump("费用类型", report["categories"])
    dump("按物业", report["by_property"])
    dump("按城市", report["by_city"])
    dump("按月趋势", report["trend"], name_key="month")

    # Detail sheet (filtered stations)
    q = _apply_filters(StationMaster.query, filters or {})
    detail = wb.create_sheet("网点明细")
    detail.append(["网点ID", "网点名称", "物业", "城市", "结算方式",
                   "结算金额", "电费单价", "服务费单价", "年度场地费", "押金", "分成比例"])
    for st in q.all():
        detail.append([
            st.site_id, st.site_name, st.property_name, st.city,
            st.settle_method, st.settle_amount, st.elec_price,
            st.service_price, st.annual_site_fee, st.deposit, st.share_ratio,
        ])
    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(bold=True)
    return wb
